from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUTPUT = Path(__file__).with_name("Plan_de_Trabajo_MV26020.xlsx")
PROJECT = "MV26020 - Automatización de Listas de Precio"
SCOPE = (
    "Diseño e implementación de una interfaz web para gestionar cambios de listas "
    "de precio mediante simulación, aprobación, ejecución controlada y trazabilidad."
)

# Los responsables y fechas definitivos no fueron proporcionados. Se usan placeholders
# y una secuencia de días hábiles para que el archivo sea transcribible y editable.
raw_tasks = [
    ("1. Incepción y definición", "Validar requisitos, alcance y restricciones del MVP", "Por definir", 100, 16),
    ("1. Incepción y definición", "Revisar historias de usuario, roles y criterios de aceptación", "Por definir", 100, 16),
    ("1. Incepción y definición", "Completar diseño de aplicación y contratos de integración", "Por definir", 100, 24),
    ("2. Diseño funcional y técnico", "Diseñar flujos de gestión global, carga masiva y cambio individual", "Por definir", 100, 32),
    ("2. Diseño funcional y técnico", "Definir modelo de estados, aprobaciones, reintentos e idempotencia", "Por definir", 100, 24),
    ("2. Diseño funcional y técnico", "Definir arquitectura, seguridad, observabilidad e infraestructura", "Por definir", 100, 32),
    ("3. Desarrollo de gestión global", "Implementar borrador, validaciones y simulación de cambio global", "Por definir", 100, 40),
    ("3. Desarrollo de gestión global", "Implementar envío a aprobación y persistencia de evidencia", "Por definir", 100, 24),
    ("4. Desarrollo de carga masiva", "Implementar descarga de plantilla y carga controlada de Excel", "Por definir", 100, 32),
    ("4. Desarrollo de carga masiva", "Implementar validación por fila, reporte de errores y resumen de impacto", "Por definir", 100, 40),
    ("5. Desarrollo de gestión individual", "Implementar búsqueda autorizada de artículo/lista y cálculo en tiempo real", "Por definir", 100, 32),
    ("5. Desarrollo de gestión individual", "Implementar envío y ejecución autorizada de cambios individuales", "Por definir", 100, 24),
    ("6. Aprobaciones, ejecución y auditoría", "Implementar bandeja de aprobaciones y controles de segregación", "Por definir", 100, 40),
    ("6. Aprobaciones, ejecución y auditoría", "Implementar ejecución posterior a aprobación y resultados correlacionados", "Por definir", 100, 40),
    ("6. Aprobaciones, ejecución y auditoría", "Implementar bitácora, consultas de historial y control por ámbito", "Por definir", 100, 32),
    ("7. Calidad, pruebas y despliegue", "Ejecutar pruebas unitarias, funcionales, seguridad e integración", "Por definir", 100, 40),
    ("7. Calidad, pruebas y despliegue", "Ejecutar pruebas basadas en propiedades para factores e idempotencia", "Por definir", 100, 24),
    ("7. Calidad, pruebas y despliegue", "Preparar evidencias de despliegue, monitoreo y rollback", "Por definir", 100, 24),
    ("8. Acta de Cierre", "Elaborar y revisar el acta de cierre del proyecto", "Por definir", 100, 8),
    ("8. Acta de Cierre", "Realizar sesión de firma de aceptación con stakeholders", "Por definir", 100, 4),
    ("8. Acta de Cierre", "Documentar lecciones aprendidas", "Por definir", 100, 8),
    ("8. Acta de Cierre", "Realizar handover y entrega formal de artefactos", "Por definir", 100, 8),
]

NAVY = "1B365D"
BLUE = "D9EAF7"
LIGHT_BLUE = "EEF5FB"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
DARK = "1F1F1F"
BORDER = Side(style="thin", color="B7C9D6")
THIN_BORDER = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)


def business_day_add(start: date, days: int) -> date:
    current = start
    remaining = days
    while remaining > 0:
        current += timedelta(days=1)
        if current.weekday() < 5:
            remaining -= 1
    return current


def next_business_day(value: date) -> date:
    while value.weekday() >= 5:
        value += timedelta(days=1)
    return value


def build_schedule() -> list[tuple]:
    cursor = date(2026, 8, 17)
    planned = []
    for deliverable, task, assignee, allocation, hours in raw_tasks:
        duration = max(1, (hours + 7) // 8)
        start = next_business_day(cursor)
        end = business_day_add(start, duration - 1)
        planned.append((deliverable, task, start, end, assignee, allocation / 100, hours))
        cursor = business_day_add(end, 1)
    return planned


def apply_header(cell):
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Arial", bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def apply_body(cell, fill_color: str, horizontal: str = "left"):
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.font = Font(name="Arial", color=DARK)
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan de Trabajo"
    summary = wb.create_sheet("Resumen de Entregables")
    schedule = build_schedule()

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.merge_cells("A1:G1")
    ws["A1"] = "PLAN DE TRABAJO"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    metadata = [
        ("Proyecto", PROJECT),
        ("Objetivo / Alcance", SCOPE),
        ("Fecha global de inicio", "Por definir"),
        ("Fecha global de fin", "Por definir"),
        ("Equipo / colaboradores", "Por definir"),
    ]
    for row, (label, value) in enumerate(metadata, start=2):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row=row, column=1).font = Font(name="Arial", bold=True, color=NAVY)
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = Font(name="Arial", color=DARK)
        value_cell.alignment = Alignment(vertical="center", wrap_text=True)
        value_cell.border = THIN_BORDER
        for col in range(3, 8):
            ws.cell(row=row, column=col).border = THIN_BORDER
    ws.row_dimensions[3].height = 34

    headers = ["Entregable", "Tarea", "Fecha Inicio", "Fecha Fin", "Colaborador Asignado", "% Asignación", "HH Esfuerzo"]
    header_row = 8
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        apply_header(cell)
    ws.row_dimensions[header_row].height = 30

    first_data_row = header_row + 1
    for index, (deliverable, task, start, end, assignee, allocation, hours) in enumerate(schedule, start=first_data_row):
        fill = LIGHT_BLUE if (index - first_data_row) % 2 == 0 else LIGHT_GRAY
        values = [deliverable, task, start, end, assignee, allocation, hours]
        for column, value in enumerate(values, start=1):
            alignment = "left"
            if column in (3, 4):
                alignment = "center"
            elif column in (6, 7):
                alignment = "right"
            cell = ws.cell(row=index, column=column, value=value)
            apply_body(cell, fill, alignment)
        ws.cell(row=index, column=3).number_format = "dd/mm/yyyy"
        ws.cell(row=index, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=index, column=6).number_format = "0%"
        ws.cell(row=index, column=7).number_format = '#,##0.00'
        ws.row_dimensions[index].height = 30

    last_data_row = first_data_row + len(schedule) - 1
    table = Table(displayName="TablaPlanTrabajo", ref=f"A{header_row}:G{last_data_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)

    total_row = last_data_row + 2
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
    total_label = ws.cell(row=total_row, column=1, value="TOTAL HH DEL PROYECTO")
    total_label.fill = PatternFill("solid", fgColor=NAVY)
    total_label.font = Font(name="Arial", bold=True, color=WHITE)
    total_label.alignment = Alignment(horizontal="right")
    for col in range(1, 7):
        ws.cell(row=total_row, column=col).border = THIN_BORDER
        ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    total_value = ws.cell(row=total_row, column=7, value=f"=SUM(G{first_data_row}:G{last_data_row})")
    total_value.fill = PatternFill("solid", fgColor=NAVY)
    total_value.font = Font(name="Arial", bold=True, color=WHITE)
    total_value.alignment = Alignment(horizontal="right")
    total_value.number_format = '#,##0.00'
    total_value.border = THIN_BORDER

    widths = {"A": 34, "B": 66, "C": 15, "D": 15, "E": 25, "F": 14, "G": 14}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.auto_filter.ref = f"A{header_row}:G{last_data_row}"

    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A5"
    summary.merge_cells("A1:E1")
    summary["A1"] = "RESUMEN DE ENTREGABLES"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 28
    summary["A2"] = "Proyecto"
    summary["B2"] = PROJECT
    summary["A2"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A2"].font = Font(name="Arial", bold=True, color=NAVY)
    summary["B2"].font = Font(name="Arial", color=DARK)
    summary.merge_cells("B2:E2")

    summary_headers = ["Entregable", "Cantidad de Tareas", "Total HH", "Fecha Inicio", "Fecha Fin"]
    for col, header in enumerate(summary_headers, start=1):
        apply_header(summary.cell(row=4, column=col, value=header))

    deliverables = []
    for item in schedule:
        if item[0] not in deliverables:
            deliverables.append(item[0])
    for idx, deliverable in enumerate(deliverables, start=5):
        fill = LIGHT_BLUE if (idx - 5) % 2 == 0 else LIGHT_GRAY
        summary.cell(idx, 1, deliverable)
        summary.cell(idx, 2, f'=COUNTIF(\'Plan de Trabajo\'!$A${first_data_row}:$A${last_data_row},A{idx})')
        summary.cell(idx, 3, f'=SUMIF(\'Plan de Trabajo\'!$A${first_data_row}:$A${last_data_row},A{idx},\'Plan de Trabajo\'!$G${first_data_row}:$G${last_data_row})')
        summary.cell(idx, 4, f'=MINIFS(\'Plan de Trabajo\'!$C${first_data_row}:$C${last_data_row},\'Plan de Trabajo\'!$A${first_data_row}:$A${last_data_row},A{idx})')
        summary.cell(idx, 5, f'=MAXIFS(\'Plan de Trabajo\'!$D${first_data_row}:$D${last_data_row},\'Plan de Trabajo\'!$A${first_data_row}:$A${last_data_row},A{idx})')
        for col in range(1, 6):
            alignment = "left" if col == 1 else ("center" if col in (2, 4, 5) else "right")
            apply_body(summary.cell(idx, col), fill, alignment)
        summary.cell(idx, 3).number_format = '#,##0.00'
        summary.cell(idx, 4).number_format = "dd/mm/yyyy"
        summary.cell(idx, 5).number_format = "dd/mm/yyyy"
        summary.row_dimensions[idx].height = 24

    summary_last_row = 4 + len(deliverables)
    summary_table = Table(displayName="TablaResumenEntregables", ref=f"A4:E{summary_last_row}")
    summary_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False)
    summary.add_table(summary_table)
    for column, width in {"A": 42, "B": 20, "C": 16, "D": 16, "E": 16}.items():
        summary.column_dimensions[column].width = width

    summary.conditional_formatting.add(
        f"C5:C{summary_last_row}",
        FormulaRule(formula=["C5>0"], fill=PatternFill("solid", fgColor="E2F0D9")),
    )

    for sheet in (ws, summary):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.sheet_view.zoomScale = 90

    wb.save(OUTPUT)

    # Verificación de estructura, fórmulas y hoja de cierre obligatoria.
    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Plan de Trabajo", "Resumen de Entregables"]
    check_plan = check["Plan de Trabajo"]
    check_summary = check["Resumen de Entregables"]
    assert check_plan.cell(total_row, 7).value == f"=SUM(G{first_data_row}:G{last_data_row})"
    assert any(check_plan.cell(row, 1).value == "8. Acta de Cierre" for row in range(first_data_row, last_data_row + 1))
    assert check_plan.cell(last_data_row, 1).value == "8. Acta de Cierre"
    assert check_summary.cell(5, 2).value.startswith("=COUNTIF(")
    assert check_summary.cell(5, 3).value.startswith("=SUMIF(")
    assert check_summary.cell(5, 4).value.startswith("=MINIFS(")
    assert check_summary.cell(5, 5).value.startswith("=MAXIFS(")
    print(f"Archivo generado y validado: {OUTPUT}")


if __name__ == "__main__":
    main()
